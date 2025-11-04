"""
Excel export utility for Seca Product Monitor
Adapted from original save_to_excel() method
"""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_to_excel(listings_queryset):
    """
    Export dealer product listings to Excel file with custom formatting

    Args:
        listings_queryset: QuerySet of DealerProductListing objects

    Returns:
        BytesIO object containing Excel file
    """
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Product Listings"

    # Define column headers with specified order
    headers = [
        'Dealer_Name',
        'Your_SKU',
        'Dealer_SKU',
        'Dealer_Listing_Name',
        'Dealer_Listing_Description',
        'Dealer_Listing_Specifications',
        'Date_Last_Scraped',
        'Diff',
        'Diff_Content_Change',
        'Product_URL',
    ]

    # Header styling
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    border_side = Side(style='thin', color='000000')
    border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Write headers
    for col_num, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border

    # Set column widths
    column_widths = {
        'A': 20,  # Dealer_Name
        'B': 15,  # Your_SKU
        'C': 15,  # Dealer_SKU
        'D': 40,  # Dealer_Listing_Name
        'E': 60,  # Dealer_Listing_Description
        'F': 60,  # Dealer_Listing_Specifications
        'G': 20,  # Date_Last_Scraped
        'H': 10,  # Diff
        'I': 40,  # Diff_Content_Change
        'J': 50,  # Product_URL
    }

    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Data styling
    data_alignment = Alignment(vertical='top', wrap_text=True)

    # Write data rows
    row_num = 2
    for listing in listings_queryset:
        # Prepare row data
        row_data = [
            listing.dealer.name,
            listing.get_your_sku(),
            listing.dealer_sku,
            listing.product_name,
            listing.description_features,
            listing.specifications,
            listing.scrape_timestamp.strftime('%Y-%m-%d %H:%M:%S') if listing.scrape_timestamp else '',
            'Yes' if listing.content_changed else 'None',
            listing.diff_content if listing.content_changed else '',
            listing.product_url,
        ]

        # Write row
        for col_num, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value if value else ''
            cell.alignment = data_alignment
            cell.border = border

            # Highlight changed products
            if col_num == 8 and listing.content_changed:  # Diff column
                cell.fill = PatternFill(start_color='FFF4CC', end_color='FFF4CC', fill_type='solid')
                cell.font = Font(bold=True, color='CC6600')

        row_num += 1

    # Freeze header row
    ws.freeze_panes = 'A2'

    # Add auto-filter
    ws.auto_filter.ref = f'A1:J{row_num - 1}'

    # Save to BytesIO
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    return excel_file


def export_dealer_to_excel(dealer, listings_queryset):
    """
    Export a single dealer's listings to Excel

    Args:
        dealer: Dealer model instance
        listings_queryset: QuerySet of DealerProductListing objects for this dealer

    Returns:
        BytesIO object containing Excel file
    """
    # Reuse the main export function
    return export_to_excel(listings_queryset)
